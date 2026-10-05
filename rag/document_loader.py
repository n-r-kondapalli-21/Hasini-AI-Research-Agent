"""
Document loader module for RAG Knowledge System.

Supports loading and extracting text content from PDF, TXT, DOCX, Markdown, and URLs.
"""

import os
import re
import logging
from pathlib import Path
from typing import Dict, Any
from urllib.parse import urlparse

logger = logging.getLogger("hasini.rag.document_loader")


class RAGDocumentLoader:
    """
    Handles text extraction from various document formats and web URLs.
    """

    SUPPORTED_EXTENSIONS = {
        ".txt": "txt",
        ".md": "markdown",
        ".markdown": "markdown",
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "docx",
    }

    @classmethod
    def load(cls, source: str) -> Dict[str, Any]:
        """
        Load text and metadata from a file path or URL.

        Args:
            source: File path (e.g., 'data/report.pdf') or URL ('https://example.com/article')

        Returns:
            Dict containing 'content', 'source', 'filename', and 'file_type'.
        """
        if not isinstance(source, str) or not source.strip():
            raise ValueError("Source must be a non-empty string.")

        source = source.strip()

        # Check if source is a URL
        parsed = urlparse(source)
        if parsed.scheme in ("http", "https"):
            return cls._load_url(source)

        # Local file handling
        return cls._load_file(source)

    @classmethod
    def _load_file(cls, filepath: str) -> Dict[str, Any]:
        """Load text from a local file path."""
        path = Path(filepath).resolve()

        if not path.exists():
            raise FileNotFoundError(f"File not found: {filepath}")

        if not path.is_file():
            raise ValueError(f"Path is not a file: {filepath}")

        ext = path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format '{ext}'. "
                f"Supported extensions: {', '.join(cls.SUPPORTED_EXTENSIONS.keys())}"
            )

        file_type = cls.SUPPORTED_EXTENSIONS[ext]
        filename = path.name

        logger.info("Loading document: %s (type: %s)", filename, file_type)

        if file_type in ("txt", "markdown"):
            content = cls._extract_plain_text(path)
        elif file_type == "pdf":
            content = cls._extract_pdf(path)
        elif file_type == "docx":
            content = cls._extract_docx(path)
        else:
            raise ValueError(f"Unhandled file type: {file_type}")

        content = content.strip()
        if not content:
            raise ValueError(f"No text content could be extracted from: {filepath}")

        return {
            "content": content,
            "source": str(path),
            "filename": filename,
            "file_type": file_type,
        }

    @classmethod
    def _extract_plain_text(cls, path: Path) -> str:
        """Extract text from TXT or Markdown files with encoding fallback."""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                with open(path, "r", encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue

        raise ValueError(f"Unable to decode text file: {path}")

    @classmethod
    def _extract_pdf(cls, path: Path) -> str:
        """Extract text from PDF file using pypdf."""
        try:
            import pypdf
        except ImportError:
            raise ImportError("pypdf package is required for PDF support. Install it via pip install pypdf.")

        text_parts = []
        try:
            reader = pypdf.PdfReader(str(path))
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n\n".join(text_parts)
        except Exception as exc:
            logger.error("Failed to extract text from PDF %s: %s", path.name, exc)
            raise ValueError(f"Error reading PDF file '{path.name}': {exc}") from exc

    @classmethod
    def _extract_docx(cls, path: Path) -> str:
        """Extract text from DOCX file using python-docx."""
        try:
            import docx
        except ImportError:
            raise ImportError("python-docx package is required for DOCX support. Install it via pip install python-docx.")

        try:
            doc = docx.Document(str(path))
            text_parts = []
            
            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            return "\n\n".join(text_parts)
        except Exception as exc:
            logger.error("Failed to extract text from DOCX %s: %s", path.name, exc)
            raise ValueError(f"Error reading DOCX file '{path.name}': {exc}") from exc

    @classmethod
    def _load_url(cls, url: str) -> Dict[str, Any]:
        """Fetch and extract text from a web URL using requests and BeautifulSoup."""
        try:
            import requests
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError("requests and beautifulsoup4 are required for URL loading.")

        logger.info("Fetching content from URL: %s", url)

        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36 HasiniAgent/1.0"
            )
        }

        try:
            response = requests.get(url, headers=headers, timeout=15)
            response.raise_for_status()

            soup = BeautifulSoup(response.content, "html.parser")

            # Remove unwanted tags
            for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
                tag.decompose()

            # Try finding main content
            main_content = soup.find("main") or soup.find("article") or soup.body or soup

            text = main_content.get_text(separator="\n")

            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            clean_text = "\n".join(chunk for chunk in chunks if chunk)

            if not clean_text:
                raise ValueError(f"No readable text extracted from URL: {url}")

            parsed = urlparse(url)
            filename = f"{parsed.netloc}{parsed.path}"
            if filename.endswith("/"):
                filename = filename[:-1]

            return {
                "content": clean_text,
                "source": url,
                "filename": filename or parsed.netloc,
                "file_type": "url",
            }
        except Exception as exc:
            logger.error("Failed to load URL %s: %s", url, exc)
            raise ValueError(f"Error fetching URL '{url}': {exc}") from exc
